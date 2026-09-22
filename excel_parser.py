#!/usr/bin/env python3
"""
Excel data parser - import master data from SAP simulator Excel files
"""
import sqlite3
from pathlib import Path
import openpyxl

class ExcelParser:
    def __init__(self, db_path="sap_simulator.db"):
        self.db_path = db_path
        self.excel_dir = Path.home() / "Downloads" / "SAP MODUL"
    
    def parse_mm_materials(self):
        """Parse MM material master data"""
        wb_path = self.excel_dir / "SAP_MM_WORK_SIMULATOR.xlsx"
        if not wb_path.exists():
            print(f"File tidak ditemukan: {wb_path}")
            return
        
        wb = openpyxl.load_workbook(wb_path, data_only=True)
        # Sheet name: 04 Material Master (from read_file output)
        if '04 Material Master' not in wb.sheetnames:
            print("Sheet '04 Material Master' tidak ditemukan")
            return
        
        ws = wb['04 Material Master']
        
        conn = sqlite3.connect(self.db_path)
        cur = conn.cursor()
        
        materials = []
        # Row 73 onwards based on read_file output
        for row in ws.iter_rows(min_row=73, max_row=100, values_only=True):
            if row[0]:  # Material ID exists
                materials.append((
                    row[0],  # Material
                    row[1],  # Description
                    row[2],  # Material Group
                    row[3],  # UoM
                    row[4],  # Plant
                    row[5],  # MRP Type
                    row[6],  # MRP Controller
                    row[7],  # Reorder Point
                    row[8],  # Safety Stock
                ))
        
        cur.executemany("""
            INSERT OR REPLACE INTO materials 
            (material_id, description, material_group, uom, plant, 
             mrp_type, mrp_controller, reorder_point, safety_stock)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, materials)
        
        conn.commit()
        print(f"Imported {len(materials)} materials")
        conn.close()
    
    def parse_mm_vendors(self):
        """Parse MM vendor master data"""
        wb_path = self.excel_dir / "SAP_SD_WORK_SIMULATOR.xlsx"
        if not wb_path.exists():
            print(f"File tidak ditemukan: {wb_path}")
            return
        
        wb = openpyxl.load_workbook(wb_path, data_only=True)
        # Sheet: 03 Customer Master (use as vendor proxy)
        if '03 Customer Master' not in wb.sheetnames:
            print("Sheet tidak ditemukan")
            return
        
        ws = wb['03 Customer Master']
        
        conn = sqlite3.connect(self.db_path)
        cur = conn.cursor()
        
        vendors = []
        for row in ws.iter_rows(min_row=59, max_row=68, values_only=True):
            if row[0]:
                vendors.append((
                    row[0],  # Customer/Vendor ID
                    row[1],  # Name
                    row[2],  # City
                    row[3],  # Payment Terms
                    row[4],  # Credit Limit
                ))
        
        cur.executemany("""
            INSERT OR REPLACE INTO vendors 
            (vendor_id, name, city, payment_terms, credit_limit)
            VALUES (?, ?, ?, ?, ?)
        """, vendors)
        
        conn.commit()
        print(f"Imported {len(vendors)} vendors")
        conn.close()
    
    def parse_hcm_employees(self):
        """Parse HCM employee master data"""
        wb_path = self.excel_dir / "SAP_HCM_WORK_SIMULATOR.xlsx"
        if not wb_path.exists():
            print(f"File tidak ditemukan: {wb_path}")
            return
        
        wb = openpyxl.load_workbook(wb_path, data_only=True)
        if '03 Employee Master' not in wb.sheetnames:
            print("Sheet tidak ditemukan")
            return
        
        ws = wb['03 Employee Master']
        
        conn = sqlite3.connect(self.db_path)
        cur = conn.cursor()
        
        employees = []
        for row in ws.iter_rows(min_row=54, max_row=63, values_only=True):
            if row[0]:
                employees.append((
                    row[0],  # Personnel No
                    row[1],  # Name
                    row[2],  # Emp Subgroup
                    row[3],  # Area
                    row[4],  # Department
                    row[5],  # Position
                    row[6],  # Basic Salary
                ))
        
        cur.executemany("""
            INSERT OR REPLACE INTO employees 
            (personnel_no, name, emp_subgroup, area, department, position, basic_salary)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, employees)
        
        conn.commit()
        print(f"Imported {len(employees)} employees")
        conn.close()
    
    def import_all(self):
        print("Starting Excel import...")
        self.parse_mm_materials()
        self.parse_mm_vendors()
        self.parse_hcm_employees()
        print("Import complete!")

if __name__ == '__main__':
    parser = ExcelParser()
    parser.import_all()
